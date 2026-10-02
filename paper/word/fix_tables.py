"""Post-process pandoc .docx files: margins, readable table layout, figure sizes."""
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEXT_W = 7.0  # inches available (Letter, 0.75" margins)


def set_cell_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:tcW")):
        tcPr.remove(e)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(int(w * 1440)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def fix(path):
    d = Document(path)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin = s.bottom_margin = Inches(0.8)
    for t in d.tables:
        n = len(t.columns)
        if n == 0:
            continue
        first = (1.55 if n > 8 else min(2.3, max(1.2, TEXT_W * 0.28))) if n > 3 else TEXT_W / n
        rest = (TEXT_W - first) / max(n - 1, 1) if n > 1 else TEXT_W
        widths = [first] + [rest] * (n - 1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        for e in tblPr.findall(qn("w:tblW")) + tblPr.findall(qn("w:tblLayout")):
            tblPr.remove(e)
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), str(int(TEXT_W * 1440)))
        tw.set(qn("w:type"), "dxa")
        tblPr.append(tw)
        for e in tblPr.findall(qn("w:tblCellMar")):
            tblPr.remove(e)
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), "40" if n > 8 else "80")
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblPr.append(mar)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        fs = Pt(7) if n > 8 else Pt(8.5)
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                if i < len(widths):
                    set_cell_width(cell, widths[i])
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = fs
    for shp in d.inline_shapes:
        if shp.width > Inches(TEXT_W):
            ratio = Inches(TEXT_W) / shp.width
            shp.width = Inches(TEXT_W)
            shp.height = int(shp.height * ratio)
    d.save(path)
    print("fixed", path)


if __name__ == "__main__":
    for p in sys.argv[1:]:
        fix(p)
